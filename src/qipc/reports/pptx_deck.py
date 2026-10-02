"""Generates the "coverage deck" (PowerPoint): a title slide plus one
ranked-table slide per quarter, presenting the least-airtime ranking
the compiler computed. Every row on every slide comes from
`qipc.compiler.rank_categories_by_airtime`; nothing here is hand-typed.
"""

from __future__ import annotations

from datetime import date

from pptx import Presentation
from pptx.util import Inches, Pt

from qipc.compiler import FilingList, rank_categories_by_airtime


def _add_title_slide(prs: Presentation, station_name: str, call_sign: str, generated_on: date) -> None:
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Quarterly Issue Coverage: Least-Airtime Ranking"
    subtitle = slide.placeholders[1]
    subtitle.text = (
        f"{station_name} ({call_sign})\n"
        f"Generated {generated_on.isoformat()}. Synthetic data; not legal advice."
    )


def _add_ranking_slide(prs: Presentation, filing: FilingList) -> None:
    ranking = rank_categories_by_airtime(filing.category_totals)

    layout = prs.slide_layouts[5]  # title only
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = (
        f"{filing.quarter.label}: Issues Ranked by Airtime (least covered first)"
    )

    rows = len(ranking) + 1
    cols = 3
    left, top, width, height = Inches(0.6), Inches(1.5), Inches(9.0), Inches(5.0)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.cell(0, 0).text = "Rank"
    table.cell(0, 1).text = "Issue category"
    table.cell(0, 2).text = "Total minutes"

    for i, (category, total_min) in enumerate(ranking, start=1):
        table.cell(i, 0).text = str(i)
        table.cell(i, 1).text = category
        table.cell(i, 2).text = str(total_min)

    for row in table.rows:
        for cell in row.cells:
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(14)

    least_category, least_total = ranking[0]
    note = slide.shapes.add_textbox(Inches(0.6), Inches(6.6), Inches(9.0), Inches(0.6))
    note.text_frame.text = (
        f"Least-covered issue this quarter: {least_category} ({least_total} minutes)."
    )


def build_coverage_deck(filings: list[FilingList], station_name: str, call_sign: str,
                         generated_on: date, output_path: str) -> None:
    prs = Presentation()
    _add_title_slide(prs, station_name, call_sign, generated_on)
    for filing in filings:
        _add_ranking_slide(prs, filing)
    prs.save(output_path)
