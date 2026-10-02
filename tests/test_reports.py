"""Word filing document and PowerPoint coverage deck generation,
including a mutation test proving the Word document is derived from
the compiled FilingList rather than retyped."""

import os
import tempfile
from datetime import date

import docx
import pptx
import pytest

from qipc import db
from qipc.compiler import build_filing_list, rank_categories_by_airtime
from qipc.reports.docx_filing import build_filing_document
from qipc.reports.pptx_deck import build_coverage_deck
from qipc.seed_data import DEFAULT_QUARTERS, generate_all_segments


@pytest.fixture
def seeded_conn():
    fd, path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    conn = db.connect(path)
    db.init_schema(conn)
    rows = generate_all_segments(DEFAULT_QUARTERS)
    db.insert_segments(conn, rows)
    yield conn
    conn.close()
    os.remove(path)


def _doc_text(path: str) -> str:
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_filing_document_contains_deadline_and_every_segment(seeded_conn, tmp_path):
    q = DEFAULT_QUARTERS[0]
    filing = build_filing_list(seeded_conn, q)
    out_path = tmp_path / "filing.docx"
    build_filing_document(filing, "WKFT-TV", "WKFT", date(2026, 7, 1), str(out_path))

    text = _doc_text(str(out_path))
    assert filing.filing_deadline in text
    for seg in filing.all_segments:
        assert seg.narrative in text


def test_filing_document_is_derived_not_retyped(seeded_conn, tmp_path):
    """Mutation test: change one narrative in the FilingList before
    rendering and confirm the rendered document changes with it."""
    q = DEFAULT_QUARTERS[0]
    filing = build_filing_list(seeded_conn, q)
    first_category = next(iter(filing.segments_by_category))
    original_narrative = filing.segments_by_category[first_category][0].narrative
    mutated_narrative = original_narrative + " [MUTATED FOR TEST]"
    filing.segments_by_category[first_category][0].narrative = mutated_narrative

    out_path = tmp_path / "filing_mutated.docx"
    build_filing_document(filing, "WKFT-TV", "WKFT", date(2026, 7, 1), str(out_path))
    text = _doc_text(str(out_path))
    assert mutated_narrative in text


def test_coverage_deck_has_a_title_slide_and_one_ranking_slide_per_quarter(seeded_conn, tmp_path):
    filings = [build_filing_list(seeded_conn, q) for q in DEFAULT_QUARTERS]
    out_path = tmp_path / "deck.pptx"
    build_coverage_deck(filings, "WKFT-TV", "WKFT", date(2026, 7, 1), str(out_path))

    deck = pptx.Presentation(str(out_path))
    assert len(deck.slides) == 1 + len(filings)


def test_coverage_deck_ranking_slide_matches_computed_ranking(seeded_conn, tmp_path):
    filings = [build_filing_list(seeded_conn, q) for q in DEFAULT_QUARTERS]
    out_path = tmp_path / "deck2.pptx"
    build_coverage_deck(filings, "WKFT-TV", "WKFT", date(2026, 7, 1), str(out_path))

    deck = pptx.Presentation(str(out_path))
    first_quarter_slide = deck.slides[1]
    ranking = rank_categories_by_airtime(filings[0].category_totals)
    least_category, least_total = ranking[0]

    found = False
    for shape in first_quarter_slide.shapes:
        if shape.has_text_frame and least_category in shape.text_frame.text:
            found = True
    assert found, f"least-covered category {least_category!r} not found on its ranking slide"
